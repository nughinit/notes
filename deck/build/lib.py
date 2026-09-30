"""Helpers para editar o deck (python-pptx + lxml). Tudo em EMU."""
import copy, re
from lxml import etree
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'a': A, 'p': P, 'r': R}
def q(tag):
    pfx, t = tag.split(':'); return '{%s}%s' % (NS[pfx], t)

IN = 914400
def E(inches): return int(round(inches * IN))

FONT_DIR = '/root/.fonts/'
_fc = {}
def font(medium=False, size=13):
    key = (medium, size)
    if key not in _fc:
        _fc[key] = ImageFont.truetype(FONT_DIR + ('Roboto-Medium.ttf' if medium else 'Roboto-Regular.ttf'), int(size * 20))
    return _fc[key]

def text_w_pt(txt, size, medium=False):
    return font(medium, size).getlength(txt) / 20.0

def wrap_lines(txt, size, width_emu, medium=False):
    """Nº de linhas (quebra gulosa por palavra) de um parágrafo."""
    wpt = width_emu / 12700.0
    n, cur = 1, ''
    for word in txt.split(' '):
        t = (cur + ' ' + word).strip() if cur else word
        if text_w_pt(t, size, medium) <= wpt or not cur:
            cur = t
        else:
            n += 1; cur = word
    return n

LINE = 1.172  # altura de linha da Roboto (em × corpo)

# ---------------------------------------------------------------- slides
def dup_slide(prs, src):
    """Duplica `src` (shapes, fundo, imagens, notas). Retorna o novo slide (no fim)."""
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag in (q('p:sp'), q('p:pic'), q('p:graphicFrame'), q('p:cxnSp'), q('p:grpSp')):
            tree.remove(el)
    # fundo
    sbg = src._element.find('p:cSld/p:bg', NS)
    nb = new._element.find('p:cSld/p:bg', NS)
    if nb is not None: nb.getparent().remove(nb)
    if sbg is not None:
        new._element.find('p:cSld', NS).insert(0, copy.deepcopy(sbg))
    # relações (imagens etc.)
    rmap = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        if rel.is_external:
            rmap[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rmap[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree:
        if el.tag in (q('p:sp'), q('p:pic'), q('p:graphicFrame'), q('p:cxnSp'), q('p:grpSp')):
            c = copy.deepcopy(el)
            for node in c.iter():
                for k, v in list(node.attrib.items()):
                    if k.startswith('{%s}' % R) and v in rmap:
                        node.set(k, rmap[v])
            tree.append(c)
    if src.has_notes_slide:
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return new

def move_after(prs, slide, ref):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    def find(s):
        for i in ids:
            if prs.part.related_part(i.rId) is s.part: return i
    node, refnode = find(slide), find(ref)
    lst.remove(node)
    refnode.addnext(node)

def shape_by_id(slide, sid):
    for el in slide.shapes._spTree:
        c = el.find('.//p:cNvPr', NS)
        if c is not None and c.get('id') == str(sid) and el.tag != q('p:nvGrpSpPr'):
            return el
    raise KeyError(sid)

def keep_only(slide, ids):
    """Remove todos os shapes exceto ids (e o rodapé/logo/título, preservados por nome)."""
    keep = set(map(str, ids))
    for el in list(slide.shapes._spTree):
        if el.tag not in (q('p:sp'), q('p:pic'), q('p:graphicFrame'), q('p:cxnSp'), q('p:grpSp')): continue
        c = el.find('.//p:cNvPr', NS)
        if c.get('id') in keep or is_chrome(el): continue
        el.getparent().remove(el)

def is_chrome(el):
    """Título (placeholder), rodapé e SEU LOGO."""
    if el.find('.//p:ph', NS) is not None and el.find('.//p:ph', NS).get('type') == 'title':
        return True
    t = ''.join((x.text or '') for x in el.iter(q('a:t')))
    return bool(re.match(r'^\s*\d+\s+·\s+Dor torácica', t)) or t.strip() == 'SEU LOGO'

def get_xfrm(el):
    x = el.find('p:spPr/a:xfrm', NS)
    if x is None: x = el.find('p:xfrm', NS)
    return x

def set_geom(el, x=None, y=None, cx=None, cy=None):
    xf = get_xfrm(el)
    off, ext = xf.find('a:off', NS), xf.find('a:ext', NS)
    if x is not None: off.set('x', str(int(x)))
    if y is not None: off.set('y', str(int(y)))
    if cx is not None: ext.set('cx', str(int(cx)))
    if cy is not None: ext.set('cy', str(int(cy)))

def geom(el):
    xf = get_xfrm(el)
    off, ext = xf.find('a:off', NS), xf.find('a:ext', NS)
    return int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy'))

# ---------------------------------------------------------------- título / rodapé
def title_el(slide):
    for el in slide.shapes._spTree.iter(q('p:sp')):
        ph = el.find('.//p:ph', NS)
        if ph is not None and ph.get('type') == 'title': return el

def title_text(slide):
    return ''.join(t.text for t in title_el(slide).iter(q('a:t')))

def set_title_suffix(slide, suffix):
    ts = list(title_el(slide).iter(q('a:t')))
    ts[-1].text = ts[-1].text + suffix

def title_lines(slide):
    w = 10771632
    return wrap_lines(title_text(slide), 28, w, medium=True)

def renumber_footers(prs):
    for i, s in enumerate(prs.slides, 1):
        for el in s.shapes._spTree.iter(q('p:sp')):
            ts = list(el.iter(q('a:t')))
            if ts and re.match(r'^\s*\d+\s+·\s+Dor torácica', ''.join((t.text or '') for t in ts)):
                full = ''.join((t.text or '') for t in ts)
                new = re.sub(r'^\s*\d+', '%d' % i, full)
                ts[0].text = new
                for t in ts[1:]: t.text = ''

# ---------------------------------------------------------------- tabelas
def table_of(el): return el.find('.//a:tbl', NS)

def set_table_font(tbl, sz):
    for n in tbl.iter(q('a:rPr'), q('a:endParaRPr'), q('a:defRPr')):
        if n.get('sz') is not None: n.set('sz', str(sz * 100))

def keep_rows(tbl, rows):
    trs = tbl.findall('a:tr', NS)
    for i, tr in enumerate(trs):
        if i not in rows: tbl.remove(tr)

def cell_paras(tc):
    out = []
    for p in tc.findall('a:txBody/a:p', NS):
        txt = ''.join(t.text or '' for t in p.iter(q('a:t')))
        r = p.find('.//a:rPr', NS)
        medium = False
        if r is not None:
            lat = r.find('a:latin', NS)
            medium = lat is not None and 'Medium' in lat.get('typeface', '') or r.get('b') == '1'
        sz = int(r.get('sz')) / 100 if r is not None and r.get('sz') else 13
        out.append((txt, sz, medium))
    return out

def fit_table(gf, min_row=None):
    """Recalcula a altura de cada linha pelo texto e ajusta o quadro. Retorna altura total."""
    tbl = table_of(gf)
    cols = [int(g.get('w')) for g in tbl.findall('a:tblGrid/a:gridCol', NS)]
    total = 0
    for tr in tbl.findall('a:tr', NS):
        need = 0
        ci = 0
        for tc in tr.findall('a:tc', NS):
            span = int(tc.get('gridSpan', '1'))
            if tc.get('hMerge') == '1':
                ci += 1; continue
            w = sum(cols[ci:ci + span]); ci += span
            pr = tc.find('a:tcPr', NS)
            ml = int(pr.get('marL', 91440)) if pr is not None else 91440
            mr = int(pr.get('marR', 91440)) if pr is not None else 91440
            mt = int(pr.get('marT', 45720)) if pr is not None else 45720
            mb = int(pr.get('marB', 45720)) if pr is not None else 45720
            h = mt + mb
            for txt, sz, med in cell_paras(tc):
                bullet = txt
                n = wrap_lines(bullet, sz, w - ml - mr, med) if txt else 1
                h += int(n * sz * LINE * 12700)
            need = max(need, h)
        hh = max(need, min_row or 0)
        tr.set('h', str(hh)); total += hh
    xf = gf.find('p:xfrm', NS)
    xf.find('a:ext', NS).set('cy', str(total))
    return total

def scale_table_width(gf, new_w):
    tbl = table_of(gf)
    gcs = tbl.findall('a:tblGrid/a:gridCol', NS)
    old = sum(int(g.get('w')) for g in gcs)
    acc = 0
    for i, g in enumerate(gcs):
        w = int(round(int(g.get('w')) * new_w / old)) if i < len(gcs) - 1 else new_w - acc
        g.set('w', str(w)); acc += w
    gf.find('p:xfrm/a:ext', NS).set('cx', str(new_w))
